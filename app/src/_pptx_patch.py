"""v2.56.1+ monkey-patch: 在 xlsxwriter 创建嵌入工作簿时启用 nan_inf_to_errors

python-pptx (pptx/chart/xlsx.py) 创建 Workbook 时只设了 {"in_memory": True}，
未启用 nan_inf_to_errors，导致 NaN/inf 数据触发 write_number() 报错：
    'NAN/INF not supported in write_number() without nan_inf_to_errors Workbook() option'

该报错会中断 chart.replace_data() 和 shapes.add_chart()，导致嵌入工作簿写入失败，
PowerPoint 打开后"编辑数据"显示空白、图表变空。

本模块在导入时自动 patch _BaseWorkbookWriter._open_worksheet，确保所有图表
创建/替换路径都不会因 NaN/inf 报错。应用层的 _safe_num_list/_clean_num/
_sanitize_chart_data 仍保留作为双重保险（将 NaN/inf 替换为 0，显示更合理）。
"""
from contextlib import contextmanager
import pptx.chart.xlsx as _pptx_xlsx

if not getattr(_pptx_xlsx, '_nan_inf_patched', False):
    from xlsxwriter import Workbook

    @contextmanager
    def _safe_open_worksheet(self, xlsx_file):
        workbook = Workbook(xlsx_file, {"in_memory": True, "nan_inf_to_errors": True})
        worksheet = workbook.add_worksheet()
        yield workbook, worksheet
        workbook.close()

    _pptx_xlsx._BaseWorkbookWriter._open_worksheet = _safe_open_worksheet
    _pptx_xlsx._nan_inf_patched = True
